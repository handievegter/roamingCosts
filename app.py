import io
import numpy as np
import pandas as pd
import streamlit as st
from pathlib import Path

st.set_page_config(
    page_title="Roaming Costs Aggre🐊",
    page_icon="gator_icon.ico",   # You can use an emoji...
    layout="wide"
)
# Auto-show a sidebar logo if any of these assets are present locally
for _p in [
    "assets/gator_logo_512.png",
    "assets/gator_icon_512x512.png",
    "assets/gator_icon_256x256.png",
    "gator_icon_512x512.png",
    "gator_icon_256x256.png",
    "assets/gator_favicon_256.png",
    "gator_icon.png",
]:
    if Path(_p).exists():
        st.sidebar.image(_p, use_container_width=True)
        break
st.title("Roaming Costs Aggre🐊")

uploaded = st.file_uploader("Upload Excel (.xlsx) in the standard format", type=["xlsx"])
redistribute_threshold = st.number_input("Redistribution threshold (ZAR)", min_value=0.0, value=10.0, step=0.5)

EXPECTED_COLUMNS = ["MSISDN","TRANSPORTER","VEHICLE REG","CALLS ROAMING","CALLS DATA","TOTAL EXCL VAT","TOTAL"]
HEADER_ROW_INDEX = 5
NUMERIC_COLS = ["CALLS ROAMING","CALLS DATA","TOTAL EXCL VAT","TOTAL"]

# This function handles columns that might have numbers stored as strings with spaces and commas.
# It coerces such mixed-type columns into proper numeric types, replacing common formatting issues.
def coerce_mixed_numeric(s: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(s):
        return s.astype(float)
    s2 = s.astype(str).str.strip().replace({"nan": np.nan, "None": np.nan})
    m = s2.notna()
    s2.loc[m] = s2.loc[m].str.replace(" ", "", regex=False).str.replace(",", ".", regex=False)
    return pd.to_numeric(s2, errors="coerce")

# Creates two blank rows to separate groups visually in the output DataFrame.
def spacer(cols): 
    return pd.DataFrame([[np.nan]*len(cols), [np.nan]*len(cols)], columns=cols)

# Redistribute sub-threshold TOTALs (in ZAR) within each base transporter to rows above threshold.
# Works in integer cents to preserve exact sums. Small rows are set to 0; their cents are
# allocated proportionally to the big rows by original TOTAL weight, with any remainder cents
# assigned by largest fractional remainders.
def redistribute_within_transporter(df_in: pd.DataFrame, threshold_zar: float) -> pd.DataFrame:
    if df_in.empty:
        return df_in
    df = df_in.copy()
    cents = (df["TOTAL"].fillna(0).astype(float) * 100).round().astype(int)
    thr = int(round(threshold_zar * 100))
    small = cents < thr
    big = cents >= thr
    if small.sum() == 0:
        return df
    # Edge case: all rows below threshold — put the whole subtotal on one random vehicle
    if big.sum() == 0:
        total_sum = int(cents.sum())
        if total_sum == 0:
            return df
        rng = np.random.default_rng()
        pick = rng.choice(df.index.to_numpy())
        cents_new = pd.Series(0, index=df.index, dtype=int)
        cents_new.loc[pick] = total_sum
        df["TOTAL"] = (cents_new / 100.0).astype(float)
        return df
    small_sum = int(cents[small].sum())
    if small_sum == 0:
        return df
    weights = cents[big].astype(float)
    total_big = weights.sum()
    if total_big <= 0:
        return df
    raw_alloc = weights * (small_sum / total_big)
    alloc_floor = pd.Series(np.floor(raw_alloc).astype(int), index=weights.index)
    remainder = small_sum - int(alloc_floor.sum())
    if remainder > 0:
        order = np.argsort(-(raw_alloc - alloc_floor))
        idx_big = cents[big].index.to_numpy()
        give_idxs = idx_big[order[:remainder]]
        add = pd.Series(0, index=df.index)
        add.loc[give_idxs] = 1
    else:
        add = pd.Series(0, index=df.index)
    cents_new = cents.copy()
    cents_new[big] = cents[big] + alloc_floor.reindex(cents[big].index, fill_value=0).astype(int) + add[big].astype(int)
    cents_new[small] = 0
    df["TOTAL"] = (cents_new / 100.0).astype(float)
    return df

if uploaded is None:
    pass
else:
    # Load Excel file with headers on row 6 (index 5) and validate expected columns are present.
    try:
        df = pd.read_excel(uploaded, header=HEADER_ROW_INDEX, dtype={"MSISDN": str})
    except Exception as e:
        st.error(f"Could not read Excel: {e}")
        st.stop()

    df.columns = [c.strip() if isinstance(c, str) else c for c in df.columns]
    missing = [c for c in EXPECTED_COLUMNS if c not in df.columns]
    if missing:
        st.error("Missing columns: " + ", ".join(missing))
        st.stop()

    df = df[EXPECTED_COLUMNS].copy()
    for col in NUMERIC_COLS:
        df[col] = coerce_mixed_numeric(df[col])

    # Always floor TOTAL to two decimals (ZAR cents) to protect customers from rounding up charges.
    df["TOTAL"] = np.floor(df["TOTAL"].astype(float) * 100.0) / 100.0

    # Create base keys by stripping any standalone 'BUP' token anywhere in the string (capitalised),
    # tolerant of surrounding spaces, hyphens, underscores, or parentheses.
    bup_pattern = r"[\s\-_]*\(?\bBUP\b\)?"
    df["TRANSPORTER_BASE"] = (
        df["TRANSPORTER"].astype(str)
          .str.replace(bup_pattern, " ", regex=True)
          .str.strip()
          .str.replace(r"\s+", " ", regex=True)
    )
    df["VEHICLE_REG_BASE"] = (
        df["VEHICLE REG"].astype(str)
          .str.replace(bup_pattern, " ", regex=True)
          .str.strip()
          .str.replace(r"\s+", " ", regex=True)
    )

    # MSISDN column now stores the count of source rows merged into each VEHICLE REG within a transporter
    # Replace MSISDN values with a compact count of rows combined per vehicle
    grouped_df = (
        df.groupby(["TRANSPORTER_BASE", "VEHICLE_REG_BASE"], as_index=False)
          .agg({"MSISDN": "size", **{c: "sum" for c in NUMERIC_COLS}})
    )
    # Rename MSISDN to ROWS_COMBINED
    grouped_df = grouped_df.rename(columns={"MSISDN": "ROWS_COMBINED"})
    # Ensure ROWS_COMBINED is integer count for display
    grouped_df["ROWS_COMBINED"] = grouped_df["ROWS_COMBINED"].astype(int)

    # Set display columns from base keys
    grouped_df["TRANSPORTER"] = grouped_df["TRANSPORTER_BASE"]
    grouped_df["VEHICLE REG"] = grouped_df["VEHICLE_REG_BASE"]

    # Floor TOTAL after aggregation to ensure cents are conservative
    grouped_df["TOTAL"] = np.floor(grouped_df["TOTAL"].astype(float) * 100.0) / 100.0

    # Optionally redistribute sub-threshold totals within each base transporter
    if redistribute_threshold is not None:
        parts_redist = []
        for _, g in grouped_df.groupby("TRANSPORTER_BASE", as_index=False, sort=False):
            parts_redist.append(redistribute_within_transporter(g, redistribute_threshold))
        redistributed_df = pd.concat(parts_redist, ignore_index=True)
        # Store redistributed values in a new column
        grouped_df["TOTAL_REDIST"] = redistributed_df["TOTAL"]
    else:
        grouped_df["TOTAL_REDIST"] = grouped_df["TOTAL"]

    # Reorder columns to the expected schema and keep helper keys for sorting only
    cols_order = ["ROWS_COMBINED", "TRANSPORTER", "VEHICLE REG", "CALLS ROAMING", "CALLS DATA", "TOTAL EXCL VAT", "TOTAL", "TOTAL_REDIST", "TRANSPORTER_BASE", "VEHICLE_REG_BASE"]
    df_display = grouped_df[cols_order]

    # Sort aggregated rows and insert double-blank spacer between base transporters
    df_sorted = df_display.sort_values(["TRANSPORTER_BASE", "VEHICLE_REG_BASE"], kind="stable").reset_index(drop=True)

    parts = []
    for _, g in df_sorted.groupby("TRANSPORTER_BASE", sort=False):
        g_no_helpers = g.drop(columns=["TRANSPORTER_BASE", "VEHICLE_REG_BASE"]).copy()
        parts.append(g_no_helpers)
        # Build a visually obvious subtotal row per transporter (use object dtype to mix str/float)
        cols_no_helpers = g_no_helpers.columns
        subtotal = pd.Series(index=cols_no_helpers, dtype="object")
        subtotal["TRANSPORTER"] = f"{g_no_helpers['TRANSPORTER'].iloc[0]} — SUBTOTAL"
        subtotal["VEHICLE REG"] = "— SUBTOTAL —"
        for c in ["ROWS_COMBINED", "CALLS ROAMING", "CALLS DATA", "TOTAL EXCL VAT", "TOTAL", "TOTAL_REDIST"]:
            if c in g_no_helpers.columns:
                subtotal[c] = float(g_no_helpers[c].fillna(0).sum())
        parts.append(pd.DataFrame([subtotal], columns=cols_no_helpers))
        parts.append(spacer(cols_no_helpers))
    out_df = pd.concat(parts, ignore_index=True)[:-2] if parts else df_sorted.drop(columns=["TRANSPORTER_BASE", "VEHICLE_REG_BASE"])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        out_df.to_excel(writer, index=False, sheet_name="Grouped")
        # Emphasize subtotal rows in the Excel export (bold + light gray fill)
        ws = writer.sheets["Grouped"]
        from openpyxl.styles import Font, PatternFill
        # Locate column indexes (1-based) for convenience
        cols_map = {name: out_df.columns.get_indexer([name])[0] + 1 for name in out_df.columns}
        veh_col = cols_map.get("VEHICLE REG")
        if veh_col is not None:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                cell = row[veh_col - 1]
                if cell.value == "— SUBTOTAL —":
                    for c in row:
                        c.font = Font(bold=True)
                        c.fill = PatternFill(fill_type="solid", fgColor="DDDDDD")

    st.download_button("⬇️ Download", buf.getvalue(),
                       file_name="grouped_by_transporter.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")